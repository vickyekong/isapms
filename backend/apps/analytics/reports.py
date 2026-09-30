import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.accounts.services import get_setting
from apps.analytics.engine import PREDICTION_DISCLAIMER
from apps.analytics.models import Prediction
from apps.records.services import attendance_stats


REPORT_TYPES = {
    "student_performance",
    "course_performance",
    "attendance",
    "results",
    "at_risk",
    "predictions",
    "faculty_department",
}


def _filtered_results(params, user):
    from apps.records.models import Result

    query = Result.objects.select_related(
        "enrolment__student__faculty",
        "enrolment__student__department",
        "enrolment__course",
        "enrolment__session",
        "enrolment__semester",
    )
    mapping = {
        "session": "enrolment__session_id",
        "semester": "enrolment__semester_id",
        "faculty": "enrolment__student__faculty_id",
        "department": "enrolment__student__department_id",
        "level": "enrolment__student__level",
        "course": "enrolment__course_id",
        "student": "enrolment__student_id",
    }
    for param, lookup in mapping.items():
        if params.get(param):
            query = query.filter(**{lookup: params.get(param)})
    if getattr(user, "role", None) == "lecturer":
        query = query.filter(enrolment__course__assignments__lecturer__user=user)
    if getattr(user, "role", None) == "student":
        query = query.filter(enrolment__student__user=user)
    return query.distinct()


def _headers_and_rows(report_type, params, user):
    from apps.analytics.insights import at_risk_alerts

    if report_type == "results" or report_type == "student_performance":
        headers = ["Matric number", "Student", "Faculty", "Department", "Level", "Course", "Session", "Semester", "CA", "Exam", "Total", "Grade", "Grade point"]
        rows = []
        for result in _filtered_results(params, user):
            student = result.enrolment.student
            assessment = result.assessment
            rows.append(
                [
                    student.matric_number,
                    student.full_name,
                    student.faculty.name,
                    student.department.name,
                    student.level,
                    result.enrolment.course.code,
                    result.enrolment.session.name,
                    result.enrolment.semester.get_name_display(),
                    float(assessment.ca_score),
                    float(assessment.exam_score) if assessment.exam_score is not None else "",
                    float(result.total_score),
                    result.grade,
                    float(result.grade_point),
                ]
            )
        return headers, rows

    if report_type == "course_performance":
        headers = ["Course", "Title", "Students", "Average total", "Passes", "Failures"]
        grouped = {}
        for result in _filtered_results(params, user):
            bucket = grouped.setdefault(
                result.enrolment.course_id,
                {"code": result.enrolment.course.code, "title": result.enrolment.course.title, "scores": [], "fail": 0},
            )
            bucket["scores"].append(float(result.total_score))
            if result.grade == "F":
                bucket["fail"] += 1
        rows = []
        for bucket in grouped.values():
            average = round(sum(bucket["scores"]) / len(bucket["scores"]), 2) if bucket["scores"] else 0
            rows.append([bucket["code"], bucket["title"], len(bucket["scores"]), average, len(bucket["scores"]) - bucket["fail"], bucket["fail"]])
        return headers, rows

    if report_type == "attendance":
        from apps.records.services import scoped_enrolments

        headers = ["Matric number", "Student", "Course", "Present", "Absent", "Excused", "Percentage", "Below threshold"]
        rows = []
        enrolments = scoped_enrolments(user).filter(status="enrolled")
        for param, lookup in {
            "session": "session_id",
            "semester": "semester_id",
            "faculty": "student__faculty_id",
            "department": "student__department_id",
            "level": "student__level",
            "course": "course_id",
            "student": "student_id",
        }.items():
            if params.get(param):
                enrolments = enrolments.filter(**{lookup: params.get(param)})
        for enrolment in enrolments.select_related("student", "course"):
            stats = attendance_stats(enrolment)
            rows.append(
                [
                    enrolment.student.matric_number,
                    enrolment.student.full_name,
                    enrolment.course.code,
                    stats["present"],
                    stats["absent"],
                    stats["excused"],
                    stats["percentage"] if stats["percentage"] is not None else "",
                    "Yes" if stats["below_threshold"] else "No",
                ]
            )
        return headers, rows

    if report_type == "predictions":
        headers = ["Matric number", "Student", "Course", "Predicted grade", "Confidence", "Risk", "Model", "Data status", "Created"]
        query = Prediction.objects.select_related("student", "course", "model_version")
        if params.get("student"):
            query = query.filter(student_id=params.get("student"))
        if params.get("course"):
            query = query.filter(course_id=params.get("course"))
        if params.get("risk_level"):
            query = query.filter(risk_level=params.get("risk_level"))
        if getattr(user, "role", None) == "lecturer":
            query = query.filter(course__assignments__lecturer__user=user).distinct()
        if getattr(user, "role", None) == "student":
            query = query.filter(student__user=user)
        rows = [
            [
                item.student.matric_number,
                item.student.full_name,
                item.course.code if item.course else "",
                item.predicted_grade,
                item.confidence,
                item.risk_level,
                item.model_version.version,
                item.data_status,
                item.created_at.strftime("%Y-%m-%d %H:%M"),
            ]
            for item in query
        ]
        return headers, rows

    if report_type == "at_risk":
        headers = ["Matric number", "Student", "Faculty", "Department", "Level", "Course", "Risk", "Reason"]
        rows = []
        for alert in at_risk_alerts(user):
            student = alert["student"]
            if params.get("risk_level") and alert["risk_level"] != params.get("risk_level"):
                continue
            if params.get("faculty") and str(student.get("faculty_id")) != str(params.get("faculty")):
                continue
            if params.get("department") and str(student.get("department_id")) != str(params.get("department")):
                continue
            if params.get("level") and str(student["level"]) != str(params.get("level")):
                continue
            if params.get("student") and str(student["id"]) != str(params.get("student")):
                continue
            rows.append(
                [
                    student["matric_number"],
                    student["full_name"],
                    student["faculty"],
                    student["department"],
                    student["level"],
                    alert["course"]["code"] if isinstance(alert["course"], dict) else (alert["course"] or ""),
                    alert["risk_level"],
                    alert["reason"],
                ]
            )
        return headers, rows

    headers = ["Faculty", "Department", "Students", "Courses"]
    from apps.academics.models import Department

    rows = []
    departments = Department.objects.filter(is_deleted=False).select_related("faculty")
    if params.get("faculty"):
        departments = departments.filter(faculty_id=params.get("faculty"))
    for department in departments:
        rows.append(
            [
                department.faculty.name,
                department.name,
                department.students.filter(is_deleted=False).count(),
                department.courses.filter(is_deleted=False).count(),
            ]
        )
    return headers, rows


def render_report(report_type, export_format, params, user):
    headers, rows = _headers_and_rows(report_type, params, user)
    rows = [["" if cell is None else cell for cell in row] for row in rows]
    institution = get_setting("institution_name", "Nexus State University")
    title = report_type.replace("_", " ").title()
    filename = f"{report_type}.{export_format if export_format != 'excel' else 'xlsx'}"
    if export_format == "csv":
        buffer = StringIO()
        writer = csv.writer(buffer)
        writer.writerow([institution, title])
        writer.writerow(headers)
        writer.writerows(rows)
        if report_type in {"predictions", "at_risk"}:
            writer.writerow([])
            writer.writerow([PREDICTION_DISCLAIMER])
        response = HttpResponse(buffer.getvalue(), content_type="text/csv")
    elif export_format in {"excel", "xlsx"}:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title[:31]
        sheet.append([institution])
        sheet.append([title])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        if report_type in {"predictions", "at_risk"}:
            sheet.append([])
            sheet.append([PREDICTION_DISCLAIMER])
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B3D2E")
        buffer = BytesIO()
        workbook.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        filename = f"{report_type}.xlsx"
    elif export_format == "pdf":
        buffer = BytesIO()
        document = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28)
        styles = getSampleStyleSheet()
        table_data = [headers] + rows
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3D2E")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DDD8")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F5")]),
                ]
            )
        )
        story = [Paragraph(institution, styles["Title"]), Paragraph(title, styles["Heading2"]), Spacer(1, 8), table]
        if report_type in {"predictions", "at_risk"}:
            story.extend([Spacer(1, 10), Paragraph(PREDICTION_DISCLAIMER, styles["Normal"])])
        document.build(story)
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    else:
        raise ValueError("Unsupported export format.")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
